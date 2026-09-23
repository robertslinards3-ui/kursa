"""Sagatavo LU budžeta veidlapu (.xlsx) pēc LU parauga.

Palaišana (no šīs mapes):  python ekranuznemumi.py   (ja vajag jaunus ekrānšāviņus)
                           python budzets.py
Rezultāts:                 Budzeta_veidlapa_AIZPILDITA.xlsx

Cenas un avoti: ../BUDZETS.md. Summas mainām tikai šeit, tad sinhronizējam BUDZETS.md.
"""
import math
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill

HERE = Path(__file__).parent
SAMPLE = HERE / "LU_paraugs_Budzeta_veidlapa.xlsx"
OUT = HERE / "Budzeta_veidlapa_AIZPILDITA.xlsx"
SHOTS_DIR = HERE / "ekranuznemumi"  # veido ekranuznemumi.py
SHOT_WIDTH_PX = 640
ROW_PX = 20
VAT = 1.21

WRAP = Alignment(wrap_text=True, vertical="top")
LINK_FONT = Font(color="0563C1", underline="single")

# Pozīcijas: (nr, nosaukums, sniedzējs, summa ar PVN, pamatojums)
ITEMS = {
    1: [
        ("1.1.", "Claude Max 5x, 2 lietotāji × 4 mēn.", "Anthropic",
         round(90 * VAT, 2) * 2 * 4, "AI programmēšanas rīks, ar kuru 2 cilvēki 4 mēnešos uzraksta MVP."),
        ("1.2.", "Virtuālais serveris (4 vCPU, 8 GB), 4 mēn.", "OVHcloud",
         round(7.21 * VAT, 2) * 4, "Serveris, kurā darbojas MVP un glabājas pilota dati (ES)."),
        ("1.3.", ".lv domēns, 1 gads", "Hostnet SIA",
         round(11.76 * VAT, 2), "MVP adrese un e-pasts saziņai ar pilota klientiem."),
        ("1.4.", "Claude API kredīts CMR nolasīšanai", "Anthropic",
         100.00, "Maksa par ~1 500 CMR foto nolasīšanu: tests (30 CMR) un pilots."),
    ],
    2: [],  # Ekspertu konsultācijas nav vajadzīgas (jurists izņemts 2026-09-23)
    3: [],  # Mobilitāte nav vajadzīga (intervijas attālināti)
}

# Budžeta lapas rindas pēc LU parauga: pirmā pozīcijas rinda, pēdējā rinda pirms "Kopā"
CATEGORY_ROWS = {1: (6, 9), 2: (12, 15), 3: (18, 26)}

# Cenu izpētes lapas. Piegādātājs: (nosaukums, saite, cena bez PVN, piegāde, cena ar PVN,
# kopā ar PVN, kritērijs, atbilst)
SHEETS = {
    "1.1": {
        "title": "1.1 Claude Max 5x (unikalitāte)",
        "suppliers": [
            ("Anthropic, Claude Max 5x", "https://claude.com/pricing",
             "90,00/mēn.", "—", "108,90/mēn.", "871,20", "Unikalitāte", "Jā"),
        ],
        "notes": [
            ("Apraksts", "Claude Max 5x (ietver Claude Code), 2 lietotāji × 4 mēneši."),
            ("Unikalitāte", "Komanda jau strādā ar Claude Code, un projekta dokumentācija ir rakstīta šim rīkam. "
                            "Max 5x ir lētākais Claude plāns, kura limiti pietiek 2 izstrādātājiem."),
            ("Specifikācija", "AI programmēšanas aģents ar piekļuvi repozitorijam; limits ≥ 5× Pro; mēneša norēķins."),
            ("Termiņš", "1.–4. mēnesis."),
        ],
    },
    "1.2": {
        "title": "1.2 Virtuālais serveris, 4 mēneši",
        "suppliers": [
            ("OVHcloud VPS-2 (4 vCore, 8 GB, 75 GB)", "https://www.ovhcloud.com/en-ie/vps/",
             "7,21/mēn.", "—", "8,72/mēn.", "34,88", "Zemākā cena", "Jā"),
            ("Hetzner CX33 (4 vCPU, 8 GB, 80 GB) + IPv4",
             "https://docs.hetzner.com/general/infrastructure-and-availability/price-adjustment/",
             "8,49/mēn. + IPv4", "—", "10,27/mēn. + IPv4", "41,08 + IPv4", "Zemākā cena", "Nē"),
            ("DigitalOcean Basic (4 vCPU, 8 GB)", "https://www.digitalocean.com/pricing/droplets",
             "48 USD/mēn.", "—", "~49,61/mēn.", "~198,44", "Zemākā cena", "Nē"),
        ],
        "notes": [
            ("Apraksts", "VPS, kurā darbojas MVP: lietotne, datubāze, rezerves kopijas."),
            ("Specifikācija", "≥ 4 vCPU, ≥ 8 GB RAM, ≥ 75 GB SSD, ES datu centrs, publiska IPv4."),
            ("Termiņš", "1.–4. mēnesis."),
            ("Piezīme", "OVH lapā cena ir “no 7,21 €”, tāpēc pirkšanas brīdī to pārbaudām. Hetzner CX33 šobrīd nav pieejams."),
        ],
    },
    "1.3": {
        "title": "1.3 .lv domēns, 1 gads",
        "suppliers": [
            ("Hostnet SIA", "https://hostnet.lv/lv/domains/", "11,76", "—", "14,23", "14,23", "Zemākā cena", "Jā"),
            ("NIC (LU MII)", "https://www.nic.lv/lv/buj-maksa-un-norekini", "12,00", "—", "14,52", "14,52",
             "Zemākā cena", "Nē"),
            ("VPSnet", "https://www.vpsnet.com/en/services/domains/lv", "20,99", "—", "25,40", "25,40",
             "Zemākā cena", "Nē"),
        ],
        "notes": [
            ("Apraksts", ".lv domēna reģistrācija 1 gadam."),
            ("Termiņš", "1. mēnesis."),
            ("Piezīme", "Hostnet lapā nav norādīts, vai cena ir ar PVN (pieņemts: bez PVN)."),
        ],
    },
    "1.4": {
        "title": "1.4 Claude API kredīts (unikalitāte)",
        "suppliers": [
            ("Anthropic, Claude API (Opus 5: 5 $ / 25 $ par 1M tokenu)",
             "https://platform.claude.com/docs/en/about-claude/pricing",
             "82,64", "—", "100,00", "100,00", "Unikalitāte", "Jā"),
        ],
        "notes": [
            ("Apraksts", "API kredīts CMR foto nolasīšanai: ~1 500 CMR × ~0,064 $."),
            ("Unikalitāte", "Tas pats piegādātājs kā 1.1. Ar šo API 1. mēnesī testējam, cik precīzi tiek nolasīti 30 īsti CMR."),
            ("Specifikācija", "Attēlu analīze, strukturēta izvade (JSON), maksa par faktisko lietojumu."),
            ("Termiņš", "1.–4. mēnesis."),
        ],
    },
}

ITEM_COL_WIDTHS = {"A": 4, "B": 30, "C": 20, "D": 14, "E": 14, "F": 14, "G": 14, "H": 13, "I": 11}


def short_url(url):
    return url.split("//", 1)[1].split("/", 1)[0].removeprefix("www.")


def add_link(cell, url, text=None):
    cell.value = text or url
    cell.hyperlink = url
    cell.font = LINK_FONT


def fill_budget_sheet(ws):
    for cat, (first, last) in CATEGORY_ROWS.items():
        for r in range(first, last + 1):
            for col in "BCDEFGH":
                ws[f"{col}{r}"].value = None
            ws.row_dimensions[r].height = 15
        for i, (nr, name, provider, amount, reason) in enumerate(ITEMS[cat]):
            r = first + i
            ws[f"B{r}"].value = nr
            ws[f"B{r}"].number_format = "@"
            ws[f"C{r}"].value = name
            ws[f"D{r}"].value = provider
            ws[f"E{r}"].value = round(amount, 2)
            ws[f"E{r}"].number_format = "0.00"
            ws[f"F{r}"].value = f"Tirgus izpēte {nr} lapā"
            ws[f"G{r}"].value = reason
            ws[f"H{r}"].value = "ES Fondi"
            for col in "BCDEFGH":
                ws[f"{col}{r}"].alignment = WRAP
                ws[f"{col}{r}"].font = copy(ws[f"{col}{first}"].font)
            ws.row_dimensions[r].height = 30
    ws["B28"].value = "Kopējās pieteikuma izmaksas\n(nepārsniedzot maksimālo summu 3506 euro)"
    for cell in ("E10", "E16", "E27", "E28"):
        ws[cell].number_format = "0.00"
    for col, width in {"C": 30, "D": 14, "E": 12, "F": 20, "G": 40, "H": 16}.items():
        ws.column_dimensions[col].width = width

    # LU zemsvītras piezīmes (6–11): teksts paliek, bet tiek ietīts tabulas platumā
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= 34:
            ws.unmerge_cells(str(rng))
    for r in range(34, 41):
        if ws[f"B{r}"].value:
            ws.merge_cells(f"B{r}:H{r}")
            ws[f"B{r}"].alignment = WRAP
            ws.row_dimensions[r].height = 15 * math.ceil(len(str(ws[f"B{r}"].value)) / 150)


def build_item_sheet(wb, template, key, spec):
    ws = wb.copy_worksheet(template)
    ws.title = key
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row > 3:
            ws.unmerge_cells(str(rng))
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
        for c in row:
            if not isinstance(c, MergedCell):
                c.value = None
                c.hyperlink = None
    for r in range(4, ws.max_row + 1):
        ws.row_dimensions[r].height = None
    for col, width in ITEM_COL_WIDTHS.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[3].height = 90

    # Zem piegādātāju tabulas notīra parauga rāmjus un stilus
    for row in ws.iter_rows(min_row=4 + len(spec["suppliers"]), max_row=ws.max_row):
        for c in row:
            if not isinstance(c, MergedCell):
                c.border, c.font, c.fill, c.alignment = Border(), Font(), PatternFill(), Alignment()

    ws["A2"].value = spec["title"]
    for i, sup in enumerate(spec["suppliers"]):
        r = 4 + i
        ws[f"A{r}"].value = f"{i + 1}."
        ws[f"B{r}"].value = sup[0]
        add_link(ws[f"C{r}"], sup[1], short_url(sup[1]))
        for col, val in zip("DEFGHI", sup[2:]):
            ws[f"{col}{r}"].value = val
        chosen = sup[7] == "Jā"
        for col in "ABCDEFGHI":
            ws[f"{col}{r}"].alignment = WRAP
            if col != "C":
                ws[f"{col}{r}"].font = Font(bold=chosen)
        ws.row_dimensions[r].height = 30

    r = 4 + len(spec["suppliers"]) + 1
    for label, text in spec["notes"]:
        ws[f"A{r}"].value = f"{label}:"
        ws[f"A{r}"].font = Font(bold=True)
        ws[f"A{r}"].alignment = Alignment(vertical="top")
        ws.merge_cells(f"A{r}:B{r}")
        ws[f"C{r}"].value = text
        ws[f"C{r}"].alignment = WRAP
        ws.merge_cells(f"C{r}:I{r}")
        ws.row_dimensions[r].height = 30 if len(text) > 95 else 15
        r += 1

    # Pretendentu bloki kā LU paraugā: nosaukums, saite, ekrānšāviņš
    r += 1
    for i, sup in enumerate(spec["suppliers"]):
        ws[f"A{r}"].value = f"{i + 1}. pretendents: {sup[0]}"
        ws[f"A{r}"].font = Font(bold=True)
        add_link(ws[f"A{r + 1}"], sup[1], f"Links: {sup[1]}")
        shot = SHOTS_DIR / f"{key}_{i + 1}.png"
        if shot.exists():
            img = XLImage(str(shot))
            scale = SHOT_WIDTH_PX / img.width
            img.width, img.height = SHOT_WIDTH_PX, round(img.height * scale)
            ws.add_image(img, f"A{r + 2}")
            rows = math.ceil(img.height / ROW_PX) + 1
        else:
            ws[f"A{r + 2}"].value = "[IEVIETOT EKRĀNŠĀVIŅU]"
            rows = 3
        for rr in range(r + 2, r + 2 + rows):
            ws.row_dimensions[rr].height = ROW_PX * 0.75  # pikseļi → punkti
        r += 3 + rows


def main():
    wb = openpyxl.load_workbook(SAMPLE)
    fill_budget_sheet(wb["Budžeta lapa"])
    template = wb["1.1. Ārējais SSD disks"]
    for key, spec in SHEETS.items():
        build_item_sheet(wb, template, key, spec)
    for old in [s for s in wb.sheetnames if s != "Budžeta lapa" and s not in SHEETS]:
        del wb[old]
    wb.active = 0
    wb.save(OUT)
    total = sum(round(item[3], 2) for items in ITEMS.values() for item in items)
    print(f"Saglabāts: {OUT.name}; kopā {total:.2f} EUR (limits 3506)")


if __name__ == "__main__":
    main()
