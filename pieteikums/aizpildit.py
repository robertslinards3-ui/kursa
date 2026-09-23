"""Aizpilda LU pieteikuma veidlapu no PIETEIKUMS.md.

Palaišana (no šīs mapes):  python aizpildit.py
Rezultāts:                 Pieteikuma_Anketa_AIZPILDITA.docx
"""
import re
from pathlib import Path

import docx

HERE = Path(__file__).parent
TEMPLATE = HERE / "LU_veidne_Pieteikuma_Anketa.docx"
OUT = HERE / "Pieteikuma_Anketa_AIZPILDITA.docx"
SOURCE = HERE / "PIETEIKUMS.md"

# Galvenās tabulas atbilžu rindas (rindas indekss veidlapā) pēc sadaļas virsraksta
ANSWER_ROWS = {
    "Komandas nosaukums": 3,
    "Kopsavilkums un mērķi": 7,
    "Idejas aktualitātes pamatojums": 9,
    "Problēmas un risinājuma raksturojums": 11,
    "Idejas esošās attīstības stadijas un potenciāla apraksts": 13,
    "Mērķa grupas apraksts": 15,
    "Tirgus stāvokļa apraksts, norādot uz sava risinājuma priekšrocībām": 17,
    "Pielietojamo metožu, tehnoloģiju un pasākumu apraksts": 19,
}
MEMBER_ROWS = {"Komandas dalībnieks 1": 21, "Komandas dalībnieks 2": 22}


def parse_sections(text):
    """Returns {'## virsraksts': [rindas]} and {'### apakšvirsraksts': [rindas]}."""
    sections, subsections = {}, {}
    current, current_sub = None, None
    for line in text.splitlines():
        if line.startswith("> ") or line.startswith("# "):
            continue
        if line.startswith("## "):
            current, current_sub = line[3:].strip(), None
            sections[current] = []
        elif line.startswith("### "):
            current_sub = line[4:].strip()
            subsections[current_sub] = []
        elif current_sub is not None:
            subsections[current_sub].append(line)
        elif current is not None:
            sections[current].append(line)
    strip = lambda ls: [l for l in ls if l.strip()]
    return ({k: strip(v) for k, v in sections.items()},
            {k: strip(v) for k, v in subsections.items()})


def add_rich_paragraph(cell, line, first=False):
    """Pievieno rindu šūnai; **teksts** kļūst trekns, '- ' kļūst par aizzīmi."""
    indent = len(line) - len(line.lstrip(" "))
    line = line.strip()
    if line.startswith("- "):
        line = "• " + line[2:]
    p = cell.paragraphs[0] if first and not cell.paragraphs[0].text else cell.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = docx.shared.Cm(0.6 * (indent // 2 + 0))
    for i, part in enumerate(re.split(r"\*\*", line)):
        if part:
            p.add_run(part).bold = (i % 2 == 1)
    return p


def fill_cell(cell, lines):
    for i, line in enumerate(lines):
        add_rich_paragraph(cell, line, first=(i == 0))


def member_fields(lines):
    fields = {}
    for line in lines:
        m = re.match(r"- ([^:]+):\s*(.*)", line.strip())
        if m:
            fields[m.group(1).strip()] = m.group(2).strip()
    return fields


def main():
    sections, subs = parse_sections(SOURCE.read_text(encoding="utf-8"))
    d = docx.Document(TEMPLATE)
    t = d.tables[0]

    for title, row in ANSWER_ROWS.items():
        fill_cell(t.rows[row].cells[0], sections[title])

    # Izvēles rūtiņas: programmas virziens un RIS3 joma
    def tick(row, wanted):
        for opt in t.rows[row].cells[0].tables[0].rows:
            if wanted in opt.cells[1].text:
                for run in opt.cells[0].paragraphs[0].runs:
                    run.text = run.text.replace("☐", "☒")
                return
        raise ValueError(f"Nav atrasta izvēle: {wanted}")
    tick(1, sections["Programmas virziens"][0].split('"')[1])
    tick(5, sections["RIS3 joma"][0])

    # Starpdisciplinaritāte zem komandas virsraksta
    label = t.rows[20].cells[0]
    for line in sections["Starpdisciplinaritāte"]:
        add_rich_paragraph(label, line)

    # Komandas dalībnieki
    for title, row in MEMBER_ROWS.items():
        f = member_fields(sections[title])
        mt = t.rows[row].cells[0].tables[0]
        def append(r, c, key):
            mt.rows[r].cells[c].paragraphs[-1].add_run(" " + f.get(key, ""))
        append(0, 0, "Vārds, uzvārds")
        append(1, 0, "Izglītības iestāde"); append(1, 1, "Studējošā apliecības numurs")
        append(2, 0, "Mācību programmas līmenis"); append(2, 1, "Telefona numurs")
        append(3, 0, "Mācību virziens"); append(3, 1, "E-pasts")
        for r, key in [(5, "Kompetences"), (7, "Pienākumi komandā"),
                       (9, "Motivācija"), (11, "Ietekme uz profesionālo izaugsmi")]:
            fill_cell(mt.rows[r].cells[0], [f.get(key, "")])

    # Trešā dalībnieka tabula nav vajadzīga (2 cilvēku komanda)
    tr = t.rows[23]._tr
    tr.getparent().remove(tr)

    # Darba grafiks
    sched = t.rows[24].cells[0].tables[0]
    for month in range(1, 5):
        lines = subs[f"{month}. mēnesis"]
        split = next(i for i, l in enumerate(lines) if l.startswith("**Darbs:**"))
        result = [l.replace("**Rezultāts:** ", "") for l in lines[:split]]
        work = lines[split + 1:]
        fill_cell(sched.rows[month].cells[1], result)
        fill_cell(sched.rows[month].cells[2], work)

    d.save(OUT)
    print(f"Saglabāts: {OUT.name}")


if __name__ == "__main__":
    main()
